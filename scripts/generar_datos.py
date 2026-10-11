"""Genera datos/Logistica_Global.xlsx y datos/Logistica_Global.csv a partir del dataset de la guía."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from datos import COLUMNAS, FILAS

SALIDA = Path(__file__).resolve().parent.parent / "datos"
SALIDA.mkdir(exist_ok=True)

# CSV (UTF-8 con BOM para que Excel respete las tildes)
with open(SALIDA / "Logistica_Global.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(COLUMNAS)
    w.writerows(FILAS)

# Excel con una tabla llamada Logistica_Global
wb = Workbook()
ws = wb.active
ws.title = "Logistica_Global"
ws.append(COLUMNAS)
for fila in FILAS:
    ws.append(list(fila))
for celda in ws[1]:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor="2A78D6")
for fila in ws.iter_rows(min_row=2):
    fila[4].number_format = "#,##0"   # Ventas_USD (entero)
    fila[5].number_format = "#,##0"   # Costo_Envio_USD (entero)
    fila[6].number_format = "0.0"     # Dias_Entrega (decimal)
for letra, ancho in zip("ABCDEFG", [11, 10, 9, 12, 12, 17, 13]):
    ws.column_dimensions[letra].width = ancho
tabla = Table(displayName="Logistica_Global", ref=f"A1:G{len(FILAS) + 1}")
tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tabla)
ws.freeze_panes = "A2"
wb.save(SALIDA / "Logistica_Global.xlsx")
print("OK:", SALIDA)
